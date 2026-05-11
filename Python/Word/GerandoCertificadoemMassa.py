from docx import Document
from docx.shared import Pt

from openpyxl import load_workbook
import os

#Pegando o caminho do arquivo
nome_arquivo_alunos = "Alunos.xlsx"
planilhaDadosAlunos = load_workbook(nome_arquivo_alunos)

#Selecionando a sheet/planilha/aba
sheet_selecionada = planilhaDadosAlunos["Nomes"]

#for = para
for linha in range(2, len(sheet_selecionada["A"]) + 1):

    #Abre o arquivo do word
    arquivoWord = Document("Certificado2.docx")

    #Seleciona o estilo
    estilo = arquivoWord.styles["Normal"]

    #Pegando o nome do aluno quando passar na célula
    nomeAluno = sheet_selecionada['A%s' % linha].value

    #for = para
    for paragrafo in arquivoWord.paragraphs:
        #if = se
        if "@nome" in paragrafo.text:
            paragrafo.text = nomeAluno
            fonte = estilo.font
            fonte.name = "Calibri (Corpo)"
            fonte.size = Pt(24)

    #Pegando o caminho da pasta e configurando o nome do certificado
    caminhoCertificados = "/home/fernandomelo/Cuso logica Progamação/Curso-Logica-Programacao-/Python/Word/" + nomeAluno + ".docx"

    #Salvando o certificado do aluno
    arquivoWord.save(caminhoCertificados)

print("Certificados gerados com sucesso!")